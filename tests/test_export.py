"""Exported notes read back as the same sentences and words, in every format;
old delivery-mark markup is not written back; nothing is written over an
existing file or the original."""

from datetime import datetime

import numpy as np
import pytest

from palmcards import export
from palmcards.notes import load_notes, parse_text
from palmcards.revisions import from_snapshot, to_snapshot
from palmcards.session import Session

TEXT = """# Opening

Good evening, everyone. / Thank you for *being* here tonight.
We started with one question: // what if practice felt easy? [rise]

[slow] Most of us rehearse in our heads. **Very** forgiving heads. [fall]

# The idea

[fast] So we built a mirror that listens. /
"""


def same(a, b):
    assert [sec.title for sec in a.sections] == [sec.title for sec in b.sections]
    assert [(s.text, s.section, [w.text for w in s.words]) for s in a.sentences] == \
        [(s.text, s.section, [w.text for w in s.words]) for s in b.sentences]
    paragraphs = lambda n: [[s.index for s in n.sentences if s.paragraph == p] for p in sorted({s.paragraph for s in n.sentences})]
    assert paragraphs(a) == paragraphs(b)


@pytest.mark.parametrize("fmt", export.FORMATS)
def test_every_format_reads_back_the_same(tmp_path, fmt):
    notes = parse_text(TEXT, "md")
    dest = tmp_path / f"out.{fmt}"
    assert export.write(notes, dest, fmt) == []
    same(load_notes(dest), notes)


def test_old_markup_is_not_written_back(tmp_path):
    notes = parse_text(TEXT, "md")
    export.write(notes, tmp_path / "out.md", "md")
    out = (tmp_path / "out.md").read_text()
    assert "Thank you for being here tonight." in out and "Very forgiving heads." in out
    assert not any(m in out for m in ("/", "*", "[slow]", "[fast]", "[rise]", "[fall]"))
    assert load_notes(tmp_path / "out.md").warnings == []


def test_untitled_sections_round_trip_in_txt_and_are_reported_in_md(tmp_path):
    notes = parse_text("First part here.\n\n\nSecond part here.\n", "txt")
    assert len(notes.sections) == 2 and notes.sections[1].title == ""
    export.write(notes, tmp_path / "a.txt", "txt")
    same(load_notes(tmp_path / "a.txt"), notes)
    lost = export.write(notes, tmp_path / "a.md", "md")
    assert lost and "Section 2" in lost[0]
    assert [s.title for s in load_notes(tmp_path / "a.md").sections] == ["", "Section 2"]


def test_never_overwrites(tmp_path):
    notes = parse_text(TEXT, "md")
    original = tmp_path / "talk.md"
    original.write_text(TEXT)
    with pytest.raises(FileExistsError, match="already exists"):
        export.write(notes, original, "md")
    assert original.read_text() == TEXT


def test_the_cli_exports_the_chosen_revision_to_the_session(tmp_path, capsys):
    original = tmp_path / "talk.docx"
    import docx

    doc = docx.Document()
    doc.add_heading("Opening", level=1)
    doc.add_paragraph("Good evening, everyone. / Thank you for *being* here.")
    doc.save(str(original))
    session = Session.create(original, root=tmp_path / "sessions")
    session.add_take(np.zeros(800, np.float32), 8000, 1.0, datetime.now(), [(0.0, 0)])
    first = session.current_revision
    edited = from_snapshot(session.snapshot(first))
    edited.sections[0].sentences[0].text = "Good evening, all."  # a later revision
    second = session.add_revision(edited)
    session.release()
    assert export.main([str(session.dir), "--revision", first, "--format", "md"]) == 0
    out = capsys.readouterr().out
    assert "docx are not kept" in out
    path = session.dir / "exports" / f"talk-{first}.md"
    assert load_notes(path).sentences[0].text == "Good evening, everyone."
    assert export.main([str(session.dir)]) == 0  # current revision, the original's format
    back = load_notes(session.dir / "exports" / f"talk-{second}.docx")
    assert back.sentences[0].text == "Good evening, all."  # the later revision's text
    assert original.exists() and export.main([str(session.dir)]) == 1  # the same file again: refused
