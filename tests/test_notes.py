from pathlib import Path

import docx
import pytest

from palmcards.notes import load_notes, parse_sentence, parse_text, split_sentences

SAMPLES = Path(__file__).resolve().parent.parent / "samples"


def words(s):
    return [w.text for w in s.words]


def stripped(raw):
    """The sentence, and how many old delivery marks were left out of it."""
    ignored = []
    return parse_sentence(raw, ignored), ignored[0]


# --- plain text; old delivery-mark markup left out ------------------------------

def test_plain_sentence():
    s, n = stripped("Good evening, everyone.")
    assert words(s) == ["Good", "evening,", "everyone."] and n == 0
    assert s.text == "Good evening, everyone." and s.raw == "Good evening, everyone."


def test_old_pause_marks_are_left_out():
    s, n = stripped("/ One / two // three //")
    assert words(s) == ["One", "two", "three"] and s.text == "One two three" and n == 4
    s, n = stripped("question:// what/ now")
    assert words(s) == ["question:", "what", "now"] and n == 2


def test_slash_inside_word_is_text():
    s, n = stripped("Bring pens and/or pencils.")
    assert "and/or" in words(s) and n == 0


def test_old_stress_marks_are_left_out():
    s, n = stripped("Thank you for *being*, here. This is **really very** important.")
    assert s.text == "Thank you for being, here. This is really very important." and n == 2
    assert s.raw == "Thank you for *being*, here. This is **really very** important."  # as written


def test_lone_asterisk_is_text():
    s, n = stripped("Five * three is fifteen.")
    assert "*" in s.text and n == 0


def test_old_pace_and_ending_marks_are_left_out():
    s, n = stripped("[SLOW] Most of us rehearse. [Rise]")
    assert words(s) == ["Most", "of", "us", "rehearse."] and n == 2
    s, n = stripped("[slow] Hurry [fast] up.")
    assert s.text == "Hurry up." and n == 2


def test_unknown_bracket_is_text():
    s, n = stripped("As shown [citation needed] here.")
    assert "[citation" in words(s) and n == 0


def test_word_offsets_and_norm():
    s, _ = stripped("/ Don't *stop*, Now!")
    for w in s.words:
        assert s.text[w.start:w.end] == w.text
    assert [w.norm for w in s.words] == ["don't", "stop", "now"]


def test_punctuation_only_token_is_not_a_word():
    s, _ = stripped("Wait — / listen.")
    assert words(s) == ["Wait", "listen."]
    assert s.text == "Wait — listen."


def test_a_file_with_old_marks_says_how_many_were_ignored_once():
    notes = parse_text("Good *evening*. / Thank you. [rise]\n\nPlain here.", "txt")
    assert [s.text for s in notes.sentences] == ["Good evening.", "Thank you.", "Plain here."]
    (warning,) = notes.warnings
    assert warning.startswith("Ignored 3 delivery marks")
    assert parse_text("Nothing marked.").warnings == []


# --- sentence splitting ------------------------------------------------------

def test_split_basic_and_trailing_ending_mark():
    assert split_sentences("Is it? [rise] Yes. [slow] Now go!") == [
        "Is it? [rise]", "Yes.", "[slow] Now go!",
    ]


def test_split_keeps_abbreviations_and_initials():
    assert split_sentences("Dr. Smith met J. Doe, e.g. at noon. Then left.") == [
        "Dr. Smith met J. Doe, e.g. at noon.", "Then left.",
    ]


def test_split_after_closing_quote_and_stress():
    assert split_sentences('He said "stop." Then *go.* Done') == [
        'He said "stop."', "Then *go.*", "Done",
    ]


def test_split_decimal_number():
    assert split_sentences("It costs 3.50 today. Really.") == ["It costs 3.50 today.", "Really."]


# --- sections and files ------------------------------------------------------

def test_txt_sections_from_headings_and_double_blank_lines():
    text = "Intro line wraps\nonto the next line.\n\nSecond para.\n\n\nNew section.\n# Titled\nLast one."
    notes = parse_text(text, "txt")
    assert [s.title for s in notes.sections] == ["", "", "Titled"]
    assert [s.text for s in notes.sections[0].sentences] == [
        "Intro line wraps onto the next line.", "Second para.",
    ]
    assert [s.index for s in notes.sentences] == [0, 1, 2, 3]
    assert [s.section for s in notes.sentences] == [0, 0, 1, 2]


def test_paragraph_index_counts_source_paragraphs():
    text = "One. Two.\n\nThree.\n# Head\nFour. Five."
    notes = parse_text(text, "txt")
    assert [s.paragraph for s in notes.sentences] == [0, 0, 1, 2, 2]


def test_md_leaves_old_marks_out_that_markdown_would_eat():
    notes = parse_text("# Opening\n\nThank you for *being* here. [rise]\n\n- A *list* item.\n", "md")
    assert notes.sections[0].title == "Opening"
    assert [s.text for s in notes.sentences] == ["Thank you for being here.", "A list item."]


def test_empty_heading_is_dropped_with_warning():
    notes = parse_text("# Empty\n\n# Real\n\nText here.", "md")
    assert [s.title for s in notes.sections] == ["Real"]
    assert any("Empty" in w for w in notes.warnings)


def test_marks_only_sentence_is_ignored_with_warning():
    notes = parse_text("Hello. // [rise]\n\nBye.", "txt")
    # "// [rise]" has no words; the [rise] attaches to "Hello." via the splitter.
    assert [s.text for s in notes.sentences] == ["Hello.", "Bye."]


def test_load_sample_md():
    notes = load_notes(SAMPLES / "sample_notes.md")
    assert [s.title for s in notes.sections] == ["Opening", "The idea"]
    assert len(notes.sentences) == 9 and notes.warnings == []  # plain text now


def test_load_txt(tmp_path):
    p = tmp_path / "talk.txt"
    p.write_text("\ufeffHello / there. [fall]\n", encoding="utf-8")
    (s,) = load_notes(p).sentences
    assert s.text == "Hello there."


def test_load_docx_headings(tmp_path):
    d = docx.Document()
    d.add_heading("Opening", level=1)
    d.add_paragraph("Thank you for *being* here. / Truly.")
    d.add_paragraph("")
    d.add_paragraph("")
    d.add_paragraph("[slow] After a break.")
    p = tmp_path / "talk.docx"
    d.save(p)

    notes = load_notes(p)
    assert [s.title for s in notes.sections] == ["Opening", ""]
    a, b, c = notes.sentences
    assert (a.text, b.text, c.text) == ("Thank you for being here.", "Truly.", "After a break.")
    assert c.section == 1 and notes.warnings[0].startswith("Ignored 3 delivery marks")


def test_doc_rejected(tmp_path):
    p = tmp_path / "old.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ValueError, match=r"\.docx"):
        load_notes(p)


def test_unsupported_and_empty_files(tmp_path):
    with pytest.raises(ValueError, match="Unsupported"):
        load_notes(tmp_path / "x.pdf")
    empty = tmp_path / "empty.txt"
    empty.write_text("   \n\n")
    with pytest.raises(ValueError, match="No text"):
        load_notes(empty)
